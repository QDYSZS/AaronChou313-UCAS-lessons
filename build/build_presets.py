# -*- coding: utf-8 -*-
"""生成 presets/*.json + manifest.json（国科大 2026 秋四份预设）。

- 京内：由 build_data.py 产出的 courses_merged.json（秋季开班与春季计划）转换，并入 syllabus 大纲链接
- 京外/研究所/基地：由 datas/ 官网导出 xlsx 解析（列结构 = 京内 24 列格式）
记录不带 id（页面加载时统一按序赋值），带 ds/key/syl 字段。
"""
import openpyxl, json, re, os
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT_DIR = os.path.join(ROOT, 'presets')
os.makedirs(OUT_DIR, exist_ok=True)

MERGE_JSON = os.path.join(HERE, 'courses_merged.json')
SYLLABI_JSON = os.path.join(ROOT, 'syllabi', 'course_syllabi.json')

PRESETS = [
    {'id': '2026秋-京内', 'name': '京内学院 · 2026秋 / 2027春', 'file': '2026秋-京内.json', 'default': True},
    {'id': '2026秋-京外', 'name': '京外学院 · 2026秋', 'file': '2026秋-京外.json', 'default': False},
    {'id': '2026秋-研究所', 'name': '研究所 · 2026秋', 'file': '2026秋-研究所.json', 'default': False},
    {'id': '2026秋-基地', 'name': '培养基地 · 2026秋', 'file': '2026秋-基地.json', 'default': False},
]
DEFAULT_XLSX = {
    '2026秋-京外': 'datas/2026年秋季学期课表-京外学院.xlsx',
    '2026秋-研究所': 'datas/2026年秋季学期课表-研究所.xlsx',
    '2026秋-基地': 'datas/2026年秋季学期课表-基地.xlsx',
}

DAY_RE = re.compile(r'周([一二三四五六日天])\((.+)\)')
DAYMAP = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5, '六': 6, '日': 7, '天': 7}


def parse_slot_m(m):
    """'周二(5-6)' / '周六(1-3,5-7)' -> list[(day,start,end)]"""
    mm = DAY_RE.match(str(m).strip())
    if not mm:
        return []
    day = DAYMAP[mm.group(1)]
    out = []
    for part in mm.group(2).split(','):
        part = part.strip()
        if '-' in part:
            a, b = part.split('-')
            out.append((day, int(a), int(b)))
        else:
            out.append((day, int(part), int(part)))
    return out


def campus_from_code(code):
    mm = re.search(r'(?:P|M|D)[A-Z]?\d{3,4}([HYZ])', code)
    return {'H': '雁栖湖', 'Y': '玉泉路', 'Z': '中关村'}[mm.group(1)] if mm else None


def load_sc_map():
    """从 courses_merged 建立 subjectCode -> {first: {second: count}} 与 known seconds"""
    sc_map = {}
    seconds = set()
    try:
        merged = json.load(open(MERGE_JSON, encoding='utf-8'))
    except FileNotFoundError:
        return {}, set()
    for c in merged:
        sc_map.setdefault(c['subjectCode'], {}).setdefault(c['first'], {})
        sec = sc_map[c['subjectCode']][c['first']]
        sec[c['second']] = sec.get(c['second'], 0) + 1
        seconds.add(c['second'])
    return sc_map, seconds


def derive(code, sc, disc, sc_map, seconds):
    fs = sc_map.get(sc)
    if fs:
        first = max(fs, key=lambda k: sum(fs[k].values()))
        sset = fs[first]
        if disc in sset:
            return (first, disc)
        second = max(sset, key=sset.get)
        return (first, second)
    if disc:
        return (disc, disc if disc in seconds else '一级学科课程')
    return ('其他 / 自设学科', '自设课程')


def parse_xlsx(path, ds, campus_fallback):
    """解析官网课表 xlsx -> records（京外/研究所/基地）"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['sheet0']
    blocks = []
    cur = None
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(row=r, column=col).value for col in range(1, 25)]
        if all(v is None for v in vals):
            continue
        if vals[0] is not None:
            if cur:
                blocks.append(cur)
            cur = {'anchor': vals, 'slots': []}
        else:
            cur['slots'].append(vals)
    if cur:
        blocks.append(cur)
    wb.close()

    sc_map, seconds = load_sc_map()
    records = []
    for c in blocks:
        a = c['anchor']
        code = str(a[2]).strip()
        name = str(a[3]).strip()
        college = str(a[1]).strip()
        disc = str(a[7]).strip() if a[7] else ''
        h, cr = str(a[8]).strip().split('/')
        sc = code[6:12] if len(code) >= 12 else ''
        first, second = derive(code, sc, disc, sc_map, seconds)
        campus = campus_from_code(code) or campus_fallback
        slots = []
        for sr in [a] + c['slots']:
            mval, wval = sr[12], sr[11]
            if not mval:
                continue
            wtext = str(wval).strip() if wval else ''
            for (day, s, e) in parse_slot_m(mval):
                slots.append({'day': day, 'p': f'{s}-{e}', 'start': s, 'end': e, 'w': wtext})
        records.append({
            'ds': ds, 'key': f'{ds}:{code}',
            'code': code, 'name': name, 'en': str(a[4]).strip() if a[4] else '',
            'college': college, 'campus': campus, 'semester': '2026年秋季学期',
            'category': str(a[5]).strip(), 'discipline': disc, 'subjectCode': sc,
            'first': first, 'second': second,
            'level': str(a[6]).strip(),
            'hours': float(h), 'credits': float(cr),
            'capacity': a[9], 'enrolled': int(a[10]) if a[10] else 0,
            'exam': str(a[15]).strip() if a[15] else '', 'teach': str(a[14]).strip() if a[14] else '',
            'chief': str(a[16]).strip() if a[16] else '', 'chiefUnit': str(a[17]).strip() if a[17] else '',
            'main': str(a[18]).strip() if a[18] else '', 'mainUnit': str(a[19]).strip() if a[19] else '',
            'ta': str(a[20]).strip() if a[20] else '', 'taUnit': str(a[21]).strip() if a[21] else '',
            'convener': str(a[22]).strip() if a[22] else '',
            'room': str(a[13]).strip() if a[13] else '', 'slots': slots,
        })
    return records


def main():
    # ---- 京内：由 courses_merged.json 秋季与春季转换 + syl ----
    merged = json.load(open(MERGE_JSON, encoding='utf-8'))
    syl_map = {}
    if os.path.exists(SYLLABI_JSON):
        syl = json.load(open(SYLLABI_JSON, encoding='utf-8'))
        syl_map = {s['match_code']: s['syllabus_url'] for s in syl.get('syllabi', []) if s.get('match_code')}

    jingnei = []
    for c in merged:
        rec = {k: v for k, v in c.items() if k != 'id'}
        rec['semester'] = '2026年秋季学期' if c['semester'] == '秋季' else '2027年春季学期'
        rec['ds'] = '2026秋-京内'
        rec['key'] = '2026秋-京内:' + c['code']
        syl = syl_map.get(c['code']) if c['semester'] == '秋季' else None
        if syl:
            rec['syl'] = syl
        jingnei.append(rec)
    write_json('2026秋-京内.json', jingnei)

    # ---- 其余预设：解析 datas xlsx ----
    for p in PRESETS:
        if p['id'] == '2026秋-京内':
            continue
        xlsx = os.path.join(ROOT, DEFAULT_XLSX[p['id']])
        campus_fb = {'2026秋-京外': '京外', '2026秋-研究所': '研究所', '2026秋-基地': '基地'}[p['id']]
        recs = parse_xlsx(xlsx, p['id'], campus_fb)
        write_json(p['file'], recs)

    # ---- manifest ----
    manifest = []
    for p in PRESETS:
        recs = json.load(open(os.path.join(OUT_DIR, p['file']), encoding='utf-8'))
        n_syl = sum(1 for r in recs if r.get('syl'))
        n_slot = sum(1 for r in recs if r.get('slots'))
        manifest.append({**p, 'count': len(recs), 'withSchedule': n_slot, 'withSyllabus': n_syl})
    with open(os.path.join(OUT_DIR, 'manifest.json'), 'w', encoding='utf-8') as f:
        json.dump(manifest, f, ensure_ascii=False, indent=1)
    for m in manifest:
        print(f"{m['id']}: {m['count']} 门 (可排课 {m['withSchedule']}, 大纲 {m['withSyllabus']})")


def write_json(fname, records):
    with open(os.path.join(OUT_DIR, fname), 'w', encoding='utf-8') as f:
        json.dump(records, f, ensure_ascii=False, separators=(',', ':'))


if __name__ == '__main__':
    main()
